from pathlib import Path
from io import BytesIO

import pytest
from openpyxl import Workbook

from sources.db_source import DBSource
from sources.upload_source import UploadSource

ROOT = Path(__file__).resolve().parents[2]


def test_uploads_sample_csvs_and_pipe_file():
    files = [(path.name, path.read_bytes()) for path in (ROOT / "sample_data").glob("*.csv")]
    files.append(("marks_pipe.txt", (ROOT / "sample_data" / "marks_pipe.txt").read_bytes()))
    source = UploadSource(files)
    try:
        assert {"orders", "products", "students", "scores", "marks_pipe"} <= set(source.list_tables())
        assert source.dialect == "sqlite"
    finally:
        source.close()


def test_duplicate_filenames_and_query_row_cap():
    raw = b"id,value\n1,a\n2,b\n3,c\n"
    source = UploadSource([("items.csv", raw), ("items.csv", raw)])
    try:
        assert source.list_tables() == ["items", "items_2"]
        assert len(source.query("SELECT * FROM items", max_rows=2)) == 2
        assert len(source.preview("items", limit=1)) == 1
    finally:
        source.close()


def test_excel_workbook_creates_table_for_each_populated_sheet():
    workbook = Workbook()
    students = workbook.active
    students.title = "Students"
    students.append(["student_id", "name"])
    students.append([1, "Ari"])
    scores = workbook.create_sheet("Scores")
    scores.append(["student_id", "score"])
    scores.append([1, 95])
    workbook.create_sheet("Empty")
    buffer = BytesIO()
    workbook.save(buffer)

    source = UploadSource([("school.xlsx", buffer.getvalue())])
    try:
        assert source.list_tables() == ["school_students", "school_scores"]
        assert source.join_keys() == [("school_students", "student_id", "school_scores", "student_id")]
        assert source.query(
            "SELECT s.name, c.score FROM school_students s "
            "JOIN school_scores c ON s.student_id = c.student_id"
        ).iloc[0].tolist() == ["Ari", 95]
    finally:
        source.close()


def test_bundled_school_workbook_contains_joinable_sheets():
    path = ROOT / "sample_data" / "school.xlsx"
    source = UploadSource([(path.name, path.read_bytes())])
    try:
        assert source.list_tables() == ["school_students", "school_scores"]
        assert ("school_students", "student_id", "school_scores", "student_id") in source.join_keys()
    finally:
        source.close()


def test_upload_row_and_table_limits(monkeypatch):
    monkeypatch.setenv("MAX_ROWS_PER_TABLE", "2")
    with pytest.raises(ValueError, match="maximum is 2 rows per table"):
        UploadSource([("large.csv", b"id,value\n1,a\n2,b\n3,c\n")])

    monkeypatch.setenv("MAX_ROWS_PER_TABLE", "1000000")
    monkeypatch.setenv("MAX_TABLES_PER_UPLOAD", "1")
    with pytest.raises(ValueError, match="more than 1 table files"):
        UploadSource([
            ("first.csv", b"id,value\n1,a\n"),
            ("second.csv", b"id,value\n2,b\n"),
        ])


def test_result_row_limit_and_sqlite_timeout(monkeypatch):
    source = UploadSource([("small.csv", b"id,value\n1,a\n2,b\n3,c\n")])
    try:
        monkeypatch.setenv("MAX_RESULT_ROWS", "2")
        assert len(source.query("SELECT * FROM small")) == 2

        monkeypatch.setenv("SQLITE_QUERY_TIMEOUT_SECONDS", "0.01")
        with pytest.raises(Exception, match="interrupted"):
            source.query(
                "WITH RECURSIVE counter(x) AS (VALUES(0) UNION ALL "
                "SELECT x + 1 FROM counter WHERE x < 10000000) SELECT sum(x) FROM counter"
            )
        assert len(source.query("SELECT * FROM small")) == 2
    finally:
        source.close()


def test_db_source_preserves_names_and_is_read_only():
    path = ROOT / "sample_data" / "chinook.sqlite"
    source = DBSource(f"sqlite:///{path}")
    try:
        assert len(source.list_tables()) == 11
        assert "InvoiceLine" in source.list_tables()
        assert ("InvoiceLine", "InvoiceId", "Invoice", "InvoiceId") in source.join_keys()
        assert "sqlite" in source.display_name()
        assert "InvoiceId" in source.preview("InvoiceLine", limit=1).columns
        with pytest.raises(Exception):
            source.query("DELETE FROM Genre")
        assert "Genre" in source.list_tables()
    finally:
        source.close()


def test_db_source_rejects_missing_sqlite_path_without_echoing_credentials():
    with pytest.raises(ValueError) as error:
        DBSource("sqlite:////definitely/not/a/database.sqlite")
    assert "password" not in str(error.value).lower()
